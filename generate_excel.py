#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
生成低空空管系统招投标数据Excel文件
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.utils import get_column_letter

# 读取CSV数据
df = pd.read_csv('/home/user/Micksiney/低空空管系统招投标数据_2024_2025.csv')

# 创建工作簿
wb = Workbook()

# ========== Sheet1: 完整数据表 ==========
ws1 = wb.active
ws1.title = "招投标完整数据"

# 定义样式
header_font = Font(bold=True, size=11, color="FFFFFF")
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
cell_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# 写入数据
for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), 1):
    for c_idx, value in enumerate(row, 1):
        cell = ws1.cell(row=r_idx, column=c_idx, value=value)
        cell.border = thin_border
        if r_idx == 1:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
        else:
            cell.alignment = cell_alignment

# 设置列宽
column_widths = {
    'A': 6,   # 序号
    'B': 12,  # 招标日期
    'C': 50,  # 项目名称
    'D': 30,  # 采购人
    'E': 35,  # 采购内容
    'F': 15,  # 采购所属行业
    'G': 15,  # 预算金额
    'H': 15,  # 中标价格
    'I': 30,  # 中标单位
    'J': 25,  # 数据来源平台
    'K': 12,  # 省份/地区
}

for col, width in column_widths.items():
    ws1.column_dimensions[col].width = width

# 冻结首行
ws1.freeze_panes = 'A2'

# ========== Sheet2: 按省份统计 ==========
ws2 = wb.create_sheet(title="按省份统计")

province_stats = df.groupby('省份/地区').agg({
    '序号': 'count',
    '预算金额(万元)': lambda x: pd.to_numeric(x.replace('-', '0'), errors='coerce').sum()
}).reset_index()
province_stats.columns = ['省份/地区', '项目数量', '预算金额合计(万元)']
province_stats = province_stats.sort_values('项目数量', ascending=False)

for r_idx, row in enumerate(dataframe_to_rows(province_stats, index=False, header=True), 1):
    for c_idx, value in enumerate(row, 1):
        cell = ws2.cell(row=r_idx, column=c_idx, value=value)
        cell.border = thin_border
        if r_idx == 1:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
        else:
            cell.alignment = cell_alignment

ws2.column_dimensions['A'].width = 15
ws2.column_dimensions['B'].width = 12
ws2.column_dimensions['C'].width = 20

# ========== Sheet3: 按行业统计 ==========
ws3 = wb.create_sheet(title="按行业统计")

industry_stats = df.groupby('采购所属行业').agg({
    '序号': 'count',
    '预算金额(万元)': lambda x: pd.to_numeric(x.replace('-', '0'), errors='coerce').sum()
}).reset_index()
industry_stats.columns = ['采购所属行业', '项目数量', '预算金额合计(万元)']
industry_stats = industry_stats.sort_values('项目数量', ascending=False)

for r_idx, row in enumerate(dataframe_to_rows(industry_stats, index=False, header=True), 1):
    for c_idx, value in enumerate(row, 1):
        cell = ws3.cell(row=r_idx, column=c_idx, value=value)
        cell.border = thin_border
        if r_idx == 1:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
        else:
            cell.alignment = cell_alignment

ws3.column_dimensions['A'].width = 20
ws3.column_dimensions['B'].width = 12
ws3.column_dimensions['C'].width = 20

# ========== Sheet4: 按平台统计 ==========
ws4 = wb.create_sheet(title="按数据来源平台统计")

platform_stats = df.groupby('数据来源平台').agg({
    '序号': 'count'
}).reset_index()
platform_stats.columns = ['数据来源平台', '项目数量']
platform_stats = platform_stats.sort_values('项目数量', ascending=False)

for r_idx, row in enumerate(dataframe_to_rows(platform_stats, index=False, header=True), 1):
    for c_idx, value in enumerate(row, 1):
        cell = ws4.cell(row=r_idx, column=c_idx, value=value)
        cell.border = thin_border
        if r_idx == 1:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
        else:
            cell.alignment = cell_alignment

ws4.column_dimensions['A'].width = 35
ws4.column_dimensions['B'].width = 12

# ========== Sheet5: 数据说明 ==========
ws5 = wb.create_sheet(title="数据说明")

description = [
    ["中国低空空管系统招投标数据汇总（2024-2025年）"],
    [""],
    ["数据说明："],
    ["1. 数据来源：通过公开渠道搜索收集，包括中国政府采购网、各省市公共资源交易平台、民航专业工程招标投标管理系统、低空界等专业平台"],
    ["2. 数据时间范围：2024年1月至2025年6月"],
    ["3. 数据覆盖范围：全国31个省市自治区的低空空管系统相关招投标项目"],
    ["4. 预算金额：以万元为单位，'-'表示未公开或暂无数据"],
    ["5. 中标价格：以万元为单位，'-'表示尚未公布中标结果或未公开"],
    ["6. 中标单位：'待公布'表示项目正在进行中或中标结果尚未公布"],
    [""],
    ["采购类型分类说明："],
    ["• 低空空管系统：低空空域管理系统、飞行管控平台等核心系统"],
    ["• 飞行服务平台：飞行服务中心、飞行综合服务平台等"],
    ["• 低空经济基建：低空经济基础设施建设工程"],
    ["• 低空感知设备：雷达、光电、ADS-B等监视感知设备"],
    ["• 低空通信设备：VHF、数据链等通信设备"],
    ["• 低空安全系统：反无人机、智能防御等安全系统"],
    ["• 低空应用服务：无人机巡检、森林防火等应用服务"],
    ["• 低空经济规划：低空经济发展规划编制服务"],
    [""],
    ["主要数据来源平台统计："],
    ["• 中国政府采购网 (ccgp.gov.cn) - 财政部唯一指定政府采购信息发布媒体"],
    ["• 全国公共资源交易平台 (ggzy.gov.cn) - 国家发改委指导"],
    ["• 民航专业工程建设项目招标投标管理系统 (zbtb.caac.gov.cn) - 民航局官方"],
    ["• 中航材招投标交易云平台 (cabidding.com.cn) - 航空行业专业平台"],
    ["• 低空界 (dikongjie.com) - 低空经济垂直信息平台"],
    ["• 各省市公共资源交易平台 - 31个省级平台"],
    [""],
    ["数据局限性说明："],
    ["• 部分项目信息可能因网站访问限制未能完整获取"],
    ["• 涉密或非公开招标项目未纳入统计"],
    ["• 数据以公开搜索渠道获取为主，可能存在遗漏"],
    ["• 建议结合官方平台进行验证和补充"],
    [""],
    ["生成时间：2025年12月16日"],
]

for r_idx, row in enumerate(description, 1):
    cell = ws5.cell(row=r_idx, column=1, value=row[0] if row else "")
    if r_idx == 1:
        cell.font = Font(bold=True, size=14)
    elif row and row[0].startswith("数据说明") or row[0].startswith("采购类型") or row[0].startswith("主要数据来源") or row[0].startswith("数据局限性"):
        cell.font = Font(bold=True, size=11)

ws5.column_dimensions['A'].width = 100

# 保存文件
output_path = '/home/user/Micksiney/低空空管系统招投标数据_2024_2025.xlsx'
wb.save(output_path)
print(f"Excel文件已生成：{output_path}")

# 打印统计信息
print(f"\n=== 数据统计 ===")
print(f"总项目数：{len(df)}个")
print(f"覆盖省份：{df['省份/地区'].nunique()}个")
print(f"数据来源平台：{df['数据来源平台'].nunique()}个")
print(f"采购行业类型：{df['采购所属行业'].nunique()}个")

# 计算预算总额
df['预算金额数值'] = pd.to_numeric(df['预算金额(万元)'].replace('-', '0'), errors='coerce')
total_budget = df['预算金额数值'].sum()
print(f"预算金额合计：{total_budget:,.2f}万元（约{total_budget/10000:.2f}亿元）")
